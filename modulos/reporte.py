from openpyxl import Workbook
from pathlib import Path
 
 
def crear_reporte(aceptados, rechazados, colisiones):
    # Crear la carpeta de salida si no existe
    Path("salida").mkdir(exist_ok=True)
 
    # Crear un archivo Excel nuevo
    libro = Workbook()
 
    # Primera hoja: novedades aceptadas
    hoja = libro.active
    hoja.title = "Carga_Sistema"
 
    columnas = [
        "consecutivo", "tipo_novedad", "tipo_documento",
        "numero_documento", "nombre_completo", "correo_corporativo",
        "codigo_cargo", "nombre_cargo", "nivel", "centro_costo",
        "area", "perfil_sistema", "requiere_aprobacion",
        "fecha_efectiva", "telefono", "fecha_vinculacion_previa"
    ]
 
    hoja.append(columnas)
 
    for novedad in aceptados:
        fila = []
 
        for columna in columnas:
            fila.append(novedad.get(columna, ""))
 
        hoja.append(fila)
 
    # Segunda hoja: novedades rechazadas
    hoja = libro.create_sheet("Rechazados")
 
    columnas_rechazados = [
        "consecutivo", "tipo", "numero_documento", "nombre",
        "clasificacion", "campo", "motivo", "valor_recibido"
    ]
 
    hoja.append(columnas_rechazados)
 
    for novedad in rechazados:
        hoja.append([
            novedad.get(columna, "")
            for columna in columnas_rechazados
        ])
 
    # Tercera hoja: colisiones de correo
    hoja = libro.create_sheet("Colisiones_Correo")
    hoja.append(["correo_base", "correo_asignado"])
 
    for correo in colisiones:
        hoja.append([
            correo.get("correo_base", ""),
            correo.get("correo_asignado", "")
        ])
 
    # Cuarta hoja: resumen
    hoja = libro.create_sheet("Resumen")
    hoja.append(["Concepto", "Cantidad"])
    hoja.append(["Novedades recibidas", len(aceptados) + len(rechazados)])
    hoja.append(["Novedades aceptadas", len(aceptados)])
    hoja.append(["Novedades rechazadas", len(rechazados)])
 
    # Guardar el archivo
    ruta = Path("salida/carga_novedades.xlsx")
    libro.save(ruta)
    libro.close()
 
    return ruta